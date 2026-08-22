from __future__ import annotations

import json
import subprocess
from datetime import datetime
from pathlib import Path
from urllib.request import urlopen

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/dondimon/git/kazakhoil-energy")
ASSETS = ROOT / "release_notes_assets"
OUT = ROOT / "release_notes"
OUT.mkdir(exist_ok=True)
DOCX = OUT / "energopulse_release_notes_2026-08-22.docx"

BLUE = RGBColor(46, 116, 181)
DEEP = RGBColor(8, 55, 47)
GREEN = RGBColor(13, 122, 93)
LIME = RGBColor(167, 217, 75)
MUTED = RGBColor(95, 111, 107)
LIGHT_GREEN = "E7F4EF"
LIGHT_BLUE = "EAF2F8"
LIGHT_AMBER = "FFF4DA"
BLACK = RGBColor(0, 0, 0)


def run(cmd: list[str]) -> str:
    return subprocess.check_output(cmd, cwd=ROOT, text=True).strip()


def api(path: str):
    try:
        with urlopen(f"http://127.0.0.1:8000{path}", timeout=60) as response:
            return json.load(response)
    except Exception:
        return None


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = "D9E5E0") -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)
        borders.append(tag)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_font(run, size=None, color=None, bold=None, italic=None):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def para(doc, text="", style=None, size=10, color=BLACK, bold=False, italic=False, before=0, after=6, align=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    if align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text)
        set_font(r, size=size, color=color, bold=bold, italic=italic)
    return p


def heading(doc, text, level=1):
    style = f"Heading {level}"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6 if level == 1 else 4)
    run = p.add_run(text)
    set_font(run, size=16 if level == 1 else 13, color=BLUE if level < 3 else RGBColor(31, 77, 120), bold=True)
    return p


def bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.28 + 0.18 * level)
    p.paragraph_format.first_line_indent = Inches(-0.16)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    set_font(r, size=9.5, color=BLACK)
    return p


def callout(doc, title, text, fill=LIGHT_GREEN):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.35)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, "D5E8E0")
    set_cell_margins(cell, 130, 150, 130, 150)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_font(r, size=10.5, color=DEEP, bold=True)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(text)
    set_font(r2, size=9, color=RGBColor(45, 61, 57))
    para(doc, "", after=4)


def metric_table(doc, items, fill=LIGHT_GREEN):
    table = doc.add_table(rows=1, cols=len(items))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    width = 6.35 / len(items)
    for i, item in enumerate(items):
        cell = table.cell(0, i)
        cell.width = Inches(width)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, fill)
        set_cell_border(cell, "D9E5E0")
        set_cell_margins(cell, 110, 120, 110, 120)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(item["value"])
        set_font(r, size=15, color=DEEP, bold=True)
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_after = Pt(0)
        r2 = p2.add_run(item["label"])
        set_font(r2, size=7.5, color=MUTED, bold=True)
    para(doc, "", after=4)


def add_image(doc, image, caption):
    path = ASSETS / image
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(3)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(6.35))
    cap = para(doc, caption, size=8.5, color=MUTED, italic=True, after=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    return cap


def simple_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [1.3, 2.05, 2.95] if len(headers) == 3 else [6.35 / len(headers)] * len(headers)
    for i, header in enumerate(headers):
        cell = table.cell(0, i)
        cell.width = Inches(widths[i])
        set_cell_shading(cell, "EAF2F8")
        set_cell_border(cell)
        set_cell_margins(cell)
        r = cell.paragraphs[0].add_run(header)
        set_font(r, size=8.5, color=DEEP, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            cell.width = Inches(widths[i])
            set_cell_border(cell)
            set_cell_margins(cell)
            r = cell.paragraphs[0].add_run(str(value))
            set_font(r, size=8.2, color=BLACK)
    para(doc, "", after=6)
    return table


def setup_doc() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    for style_name in ("Normal", "List Bullet"):
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(10)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("ЭнергоПульс release notes · Казахойл Актобе")
    set_font(r, size=8, color=MUTED)
    return doc


def main():
    imports = api("/api/v1/imports") or []
    dashboard = api("/api/v1/dashboards/energy-business") or {}
    ai = api("/api/v1/ai/overview-forecast") or {}
    commit_log = run(["git", "log", "--since=2 days ago", "--date=short", "--pretty=format:%h | %ad | %s"])

    ready = sum(1 for item in imports if item.get("status") in ("ready_to_publish", "published"))
    total_rows = sum(int(item.get("total_rows") or 0) for item in imports)
    kinds = sorted({str(item.get("dataset_kind")) for item in imports})
    monthly = dashboard.get("monthly_series") or []
    daily = dashboard.get("daily_series") or []
    forecast = dashboard.get("forecast") or {}
    ai_basis = ai.get("data_basis") or {}

    doc = setup_doc()
    para(doc, "RELEASE NOTES", size=9, color=GREEN, bold=True, after=6)
    para(doc, "ЭнергоПульс: обновления продукта и функционал системы", size=24, color=DEEP, bold=True, after=4)
    para(doc, "Клиентская версия · 22 августа 2026", size=12, color=MUTED, after=14)
    callout(
        doc,
        "Кратко",
        "Обновление сфокусировано на более устойчивом Excel ETL, прозрачном энергобалансе, поиске по крупным таблицам, понятной диагностике недостающих данных и AI-помощнике «ЭнергоПульс AI».",
    )
    metric_table(
        doc,
        [
            {"value": str(len(imports)), "label": "загруженных файлов"},
            {"value": str(ready), "label": "готовы к расчёту"},
            {"value": f"{total_rows:,}".replace(",", " "), "label": "обработанных строк"},
            {"value": str(len(kinds)), "label": "типа источников"},
        ],
    )

    heading(doc, "1. Новые обновления", 1)
    updates = [
        "Универсальный Excel ETL: парсер ориентируется на структуру листов и индексы Excel, а не на жёстко заданные названия объектов.",
        "Названия потребителей, подстанций и строк баланса теперь берутся из Excel без преобразований и ручных словарей в бизнес-логике.",
        "Добавлена поддержка актов потребления как отдельного типа данных; техбаланс, акты и daily-файлы отображаются в журнале загрузок.",
        "ЭнергоПульс AI на странице «Сводка» переведён на следующий календарный месяц, учитывает текущую дату и не переиспользует устаревший кэш.",
        "Перед прогнозным инсайтом система запрашивает Open-Meteo за целевой период и формирует отдельный погодный комментарий.",
        "Блок «ЭнергоПульс AI» переработан в кубики: главные инсайты, погодный комментарий, expandable детали и переход к странице прогноза.",
        "Прогноз и назначение погодных регионов стали устойчивее: регион сохраняется по id и стабильному ключу имени потребителя.",
        "Сверка энергобаланса теперь объясняет, если ежедневная сводка загружена за другие месяцы, а не показывает ложное «нет данных».",
        "В разделе «Потребители» добавлен локальный поиск по потребителю, компании, подстанции и погодному региону.",
        "В таблицах объектов и потребления добавлен поиск по наименованию, номеру ПУ, подстанции, коэффициенту, расходу и служебным полям.",
        "Для общего расхода по счётчикам интерфейс показывает номер счётчика и источник данных: номер ПУ из столбца C, расход рассчитывается по формуле (G - F) × E.",
        "В «Пиках и аномалиях» список периодов ограничен месяцами, где есть рассчитанный дневной ряд с положительным значением.",
        "Раздел «Месячная сверка» заменён на диагностический экран: он показывает, какие месяцы имеют пару техбаланс + ежедневная сводка и какого файла не хватает.",
        "Пользовательский AI-раздел переименован в «ЭнергоПульс AI» без изменения технической интеграции.",
        "По данным подтверждены объекты Кожасай: линии «ПС 110/35/6 ВЛ 35кВ Кожасай №1/№2», а также узлы «П/С 35/6 \"Кожасай\"» и «П/С 35/6 \"БКНС Кожасай\"».",
    ]
    for item in updates:
        bullet(doc, item)

    heading(doc, "Скриншоты обновлений", 2)
    add_image(doc, "01_overview.png", "Сводка: компактный блок расчётов, протокол загрузки и ЭнергоПульс AI с погодным инсайтом.")
    add_image(doc, "02_energy_balance.png", "Энергобаланс: структура потребления, дневная нагрузка и ключевые KPI по последнему месяцу.")
    add_image(doc, "03_forecast.png", "Прогноз: сценарный расчёт с погодной поправкой, надёжностью и графиком факт/прогноз.")
    add_image(doc, "04_data_quality.png", "Исходные данные: журнал файлов, качество проверки и готовность к расчёту.")

    heading(doc, "2. Функционал системы", 1)
    simple_table(
        doc,
        ["Модуль", "Что делает", "Пользовательская ценность"],
        [
            ("Загрузка и проверка", "Принимает Excel-файлы нужного формата, классифицирует тип источника, считает строки и замечания.", "Пользователь видит, какие файлы готовы к расчёту и где требуется внимание."),
            ("Excel ETL", "Извлекает данные из техбаланса, актов потребления и ежедневных сводок по структуре листов.", "Меньше зависимости от конкретных названий и ручных корректировок."),
            ("Энергобаланс", "Считает общий вход, собственное потребление КОА и потребление сторонних организаций.", "Даёт управленческую картину распределения электроэнергии."),
            ("Дневная нагрузка", "Строит суточный ряд, пики и резкие изменения при наличии ежедневной сводки.", "Помогает быстро увидеть аномальные дни и отклонения."),
            ("Потребители", "Показывает внешних потребителей из Excel, поиск по списку и привязку к погодным регионам.", "Готовит основу для погодной поправки и детализации прогноза."),
            ("Прогноз", "Строит следующий период по истории, календарю, погоде и пользовательским событиям.", "Помогает оценить диапазон будущего потребления и риски."),
            ("ЭнергоПульс AI", "Сжимает расчёт, погодный контекст и качество данных в короткие инсайты.", "Руководитель получает не таблицу, а понятные действия и ограничения."),
            ("Сверка данных", "Сопоставляет месячный техбаланс с ежедневными сводками по совпадающим периодам и показывает недостающие файлы.", "Пользователь понимает, почему сверка недоступна и что именно нужно загрузить."),
            ("Пики и аномалии", "Показывает только периоды, где есть пригодный дневной ряд для расчёта.", "Снижает риск выбрать месяц без достаточных данных."),
        ],
    )

    heading(doc, "3. Текущий статус данных", 1)
    latest_period = dashboard.get("meta", {}).get("latest_period") or "не определён"
    daily_periods = sorted({str(item.get("period")) for item in daily if item.get("period")})
    monthly_periods = [str(item.get("period")) for item in monthly]
    callout(
        doc,
        "Состояние на момент сборки документа",
        f"Последний период энергобаланса: {latest_period}. Месячные периоды: {', '.join(monthly_periods) or 'нет'}. Суточные периоды: {', '.join(daily_periods) or 'нет'}. ЭнергоПульс AI ориентирован на {ai_basis.get('target_forecast_period', 'следующий месяц')}; погодный контекст: {ai_basis.get('weather_status', 'нет статуса')} ({ai_basis.get('weather_days', 0)} дней).",
        fill=LIGHT_BLUE,
    )

    heading(doc, "4. История релиза и проверки", 1)
    simple_table(
        doc,
        ["Источник", "Содержание", "Комментарий"],
        [
            ("Git commits за 2 дня", commit_log or "Коммитов не найдено", "Включены изменения последнего релизного коммита."),
            ("Текущая сессия", "Незакоммиченные изменения: поиск, источники счётчиков, пики, месячная сверка, ЭнергоПульс AI, UI и тесты.", "Документ учитывает текущий рабочий результат, а не только коммиты."),
            ("Проверки", "npm run build; targeted API unittest для daily dashboard", "Проверки проходили в ходе текущей сессии."),
        ],
    )
    heading(doc, "5. Что важно донести клиенту", 1)
    for item in [
        "Система уже работает с несколькими источниками Excel и показывает, какие данные готовы для расчётов.",
        "ЭнергоПульс AI не заменяет расчёт: он объясняет расчёт, погодный контекст и ограничения качества данных.",
        "Для точной месячной сверки ежедневные сводки должны соответствовать месяцам техбаланса.",
        "Для Кожасай в данных нужно различать уровни: 110/35/6 относится к питающей ПС Кенкияк, а Кожасай представлен линиями 35 кВ и подстанциями 35/6.",
    ]:
        bullet(doc, item)

    doc.core_properties.title = "ЭнергоПульс: release notes"
    doc.core_properties.subject = "Клиентские релиз-ноты по обновлениям и функционалу системы"
    doc.core_properties.author = "Codex"
    doc.core_properties.created = datetime(2026, 8, 22)
    doc.save(DOCX)
    print(DOCX)


if __name__ == "__main__":
    main()
