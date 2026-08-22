from io import BytesIO
import unittest

import openpyxl

from app.models import DatasetKind, ValidationSeverity
from app.services.ingestion import parse_file


def workbook_payload(sheet_name: str, rows: list[list[object]]) -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    for row in rows:
        sheet.append(row)
    payload = BytesIO()
    workbook.save(payload)
    return payload.getvalue()


def technical_layout_payload() -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Произвольное имя"
    sheet.cell(row=294, column=2, value="35 kV balance block")
    for row_index in (295, 296):
        sheet.cell(row=row_index, column=1, value=f"Line {row_index}")
        sheet.cell(row=row_index, column=3, value=f"M{row_index}")
        sheet.cell(row=row_index, column=5, value=10)
        sheet.cell(row=row_index, column=6, value=1)
        sheet.cell(row=row_index, column=7, value=2)
        sheet.cell(row=row_index, column=8, value=f"=(G{row_index}-F{row_index})*E{row_index}")
    sheet.cell(row=307, column=2, value="Substation block")
    for row_index in (308, 323, 328):
        sheet.cell(row=row_index, column=1, value=f"Feeder {row_index}")
        sheet.cell(row=row_index, column=3, value=f"M{row_index}")
        sheet.cell(row=row_index, column=5, value=10)
        sheet.cell(row=row_index, column=6, value=1)
        sheet.cell(row=row_index, column=7, value=2)
        sheet.cell(row=row_index, column=8, value=f"=(G{row_index}-F{row_index})*E{row_index}")
    payload = BytesIO()
    workbook.save(payload)
    return payload.getvalue()


def commercial_consumption_payload() -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Произвольный акт"
    headers = [
        "№ п/п",
        "Номер договора",
        "Место установки счетчика",
        "Заводской номер",
        "Показания начало",
        "Показания конец",
        "Разница",
        "Коэффициент",
        "Потребление",
        "Потребление общее",
    ]
    for offset, value in enumerate(headers, start=2):
        sheet.cell(row=4, column=offset, value=value)
    for row_index in range(5, 18):
        sheet.cell(row=row_index, column=2, value=row_index - 4)
        sheet.cell(row=row_index, column=3, value=20084)
        sheet.cell(row=row_index, column=4, value=f"Consumer {row_index}")
        sheet.cell(row=row_index, column=5, value=f"M{row_index}")
        sheet.cell(row=row_index, column=6, value=1)
        sheet.cell(row=row_index, column=7, value=2)
        sheet.cell(row=row_index, column=8, value=f"=G{row_index}-F{row_index}")
        sheet.cell(row=row_index, column=9, value=10)
        sheet.cell(row=row_index, column=10, value=f"=H{row_index}*I{row_index}")
    payload = BytesIO()
    workbook.save(payload)
    return payload.getvalue()


class IngestionValidationTests(unittest.TestCase):
    def test_unknown_workbook_template_returns_validation_error(self) -> None:
        payload = workbook_payload(
            "Комм. Июль 2026г.",
            [
                ["№ п/п", "Место установки счетчика", "Заводской номер", "Потребление"],
                [1, "ПС Север", "51555209", 100],
            ],
        )

        parsed = parse_file("7. Акт потребления за июль 2026г_.xlsx", payload)

        self.assertEqual(parsed.dataset_kind, DatasetKind.unknown)
        self.assertEqual(len(parsed.rows), 2)
        self.assertTrue(
            any(
                issue.severity == ValidationSeverity.error
                and issue.rule_code == "UNSUPPORTED_DATASET_FORMAT"
                for issue in parsed.issues
            )
        )
        self.assertTrue(any("строки 294:305" in issue.message for issue in parsed.issues))

    def test_technical_balance_detected_by_excel_layout_not_sheet_name(self) -> None:
        parsed = parse_file("technical.xlsx", technical_layout_payload())

        self.assertEqual(parsed.dataset_kind, DatasetKind.technical_balance)
        self.assertFalse([issue for issue in parsed.issues if issue.severity == ValidationSeverity.error])

    def test_commercial_consumption_detected_by_excel_layout_not_sheet_name(self) -> None:
        parsed = parse_file("act.xlsx", commercial_consumption_payload())

        self.assertEqual(parsed.dataset_kind, DatasetKind.commercial_consumption)
        self.assertFalse([issue for issue in parsed.issues if issue.severity == ValidationSeverity.error])

    def test_daily_sheet_names_still_pass_validation(self) -> None:
        payload = workbook_payload(
            "01.07",
            [
                ["Итого по вводам 6 кВ"],
                ['Яч.212 "Каспий нефть-2"', "ARTM", 51555226, None, 10, 0, 2, 20],
            ],
        )

        parsed = parse_file("Ежедневная сводка потребления июль 2026.xlsx", payload)

        self.assertEqual(parsed.dataset_kind, DatasetKind.daily_summary)
        self.assertFalse([issue for issue in parsed.issues if issue.severity == ValidationSeverity.error])


if __name__ == "__main__":
    unittest.main()
